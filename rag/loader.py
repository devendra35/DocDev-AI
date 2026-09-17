from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument


logger = logging.getLogger(__name__)


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class DocumentLoaderError(Exception):
    """Base exception for document loading errors."""


class UnsupportedFileTypeError(DocumentLoaderError):
    """Raised when a file format is not supported."""


class EmptyDocumentError(DocumentLoaderError):
    """Raised when a document contains no readable text."""


@dataclass(frozen=True)
class LoadedDocument:
    """Represents extracted document content before chunking."""

    text: str
    source: str
    file_type: str
    page: int | None = None

    @property
    def metadata(self) -> dict:
        return {
            "source": self.source,
            "file_type": self.file_type,
            "page": self.page,
        }


def validate_file(file_path: str | Path) -> Path:
    """Validate that the file exists and is supported."""

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    if not path.is_file():
        raise DocumentLoaderError(f"Path is not a file: {path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))

        raise UnsupportedFileTypeError(
            f"Unsupported file type '{extension}'. "
            f"Supported formats: {supported}"
        )

    return path


def clean_text(text: str) -> str:
    """Clean extracted text while preserving meaningful content."""

    lines = [line.strip() for line in text.splitlines()]

    cleaned_lines = [
        line
        for line in lines
        if line
    ]

    return "\n".join(cleaned_lines).strip()


def load_pdf(path: Path) -> list[LoadedDocument]:
    """Extract text from every readable PDF page."""

    documents: list[LoadedDocument] = []

    try:
        pdf = pymupdf.open(path)

        try:
            for page_number, page in enumerate(pdf, start=1):

                text = clean_text(
                    page.get_text("text")
                )

                if not text:
                    logger.warning(
                        "No readable text on page %s of %s",
                        page_number,
                        path.name,
                    )
                    continue

                documents.append(
                    LoadedDocument(
                        text=text,
                        source=path.name,
                        file_type="pdf",
                        page=page_number,
                    )
                )

        finally:
            pdf.close()

    except Exception as exc:
        raise DocumentLoaderError(
            f"Failed to read PDF '{path.name}': {exc}"
        ) from exc

    return documents


def load_docx(path: Path) -> list[LoadedDocument]:
    """Extract text from a DOCX document."""

    try:
        document = DocxDocument(path)

        paragraphs = []

        for paragraph in document.paragraphs:
            text = clean_text(paragraph.text)

            if text:
                paragraphs.append(text)

        text = "\n".join(paragraphs).strip()

        if not text:
            return []

        return [
            LoadedDocument(
                text=text,
                source=path.name,
                file_type="docx",
                page=None,
            )
        ]

    except Exception as exc:
        raise DocumentLoaderError(
            f"Failed to read DOCX '{path.name}': {exc}"
        ) from exc


def load_txt(path: Path) -> list[LoadedDocument]:
    """Extract text from a UTF-8 text file."""

    try:
        text = path.read_text(
            encoding="utf-8",
            errors="replace",
        )

        text = clean_text(text)

        if not text:
            return []

        return [
            LoadedDocument(
                text=text,
                source=path.name,
                file_type="txt",
                page=None,
            )
        ]

    except Exception as exc:
        raise DocumentLoaderError(
            f"Failed to read TXT '{path.name}': {exc}"
        ) from exc


def load_document(
    file_path: str | Path,
) -> list[LoadedDocument]:
    """
    Load a supported document.

    Supported formats:
    PDF, DOCX and TXT.
    """

    path = validate_file(file_path)

    extension = path.suffix.lower()

    logger.info(
        "Loading document: %s",
        path.name,
    )

    if extension == ".pdf":
        documents = load_pdf(path)

    elif extension == ".docx":
        documents = load_docx(path)

    elif extension == ".txt":
        documents = load_txt(path)

    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type: {extension}"
        )

    if not documents:
        raise EmptyDocumentError(
            f"No readable text found in '{path.name}'. "
            "The document may be empty or image-only."
        )

    logger.info(
        "Loaded %s text unit(s) from %s",
        len(documents),
        path.name,
    )

    return documents